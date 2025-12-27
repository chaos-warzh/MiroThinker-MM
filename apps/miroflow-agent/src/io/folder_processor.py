# Copyright 2025 Miromind.ai
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#    http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""
Folder Processor Module

This module provides functionality to process all files in a folder
and prepare them for multi-modal LLM processing.

Supports:
- Images: jpg, jpeg, png, gif, webp
- Videos: mp4, avi, mov, mkv, webm, flv, wmv, m4v
- Audio: wav, mp3, m4a
- Documents: pdf, docx, doc, txt, xlsx, xls, pptx, ppt, html, htm
- Data: json, jsonld, csv
- Archives: zip

Usage:
    from src.io.folder_processor import process_folder_for_task
    
    task_content, task_description, multimodal_files = process_folder_for_task(
        folder_path="data/000",
        query="请根据这个pdf文件和图片的内容，整理重要文献"
    )
"""

import os
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

# Import existing converters and utilities from input_handler
from .input_handler import (
    DocumentConverterResult,
    XlsxConverter,
    DocxConverter,
    HtmlConverter,
    PptxConverter,
    ZipConverter,
    process_input,
)

# Try to import optional dependencies
try:
    import pdfminer.high_level
    from pdfminer.pdfpage import PDFPage
    HAS_PDFMINER = True
except ImportError:
    HAS_PDFMINER = False

try:
    from markitdown import MarkItDown
    HAS_MARKITDOWN = True
except ImportError:
    HAS_MARKITDOWN = False

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

import json


# File type categories
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
VIDEO_EXTENSIONS = {".mp4", ".avi", ".mov", ".mkv", ".webm", ".flv", ".wmv", ".m4v"}
AUDIO_EXTENSIONS = {".wav", ".mp3", ".m4a"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".html", ".htm"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xls", ".csv"}
PRESENTATION_EXTENSIONS = {".pptx", ".ppt"}
DATA_EXTENSIONS = {".jsonld", ".json"}
ARCHIVE_EXTENSIONS = {".zip"}


@dataclass
class FileInfo:
    """Information about a single file."""
    path: str
    name: str
    extension: str
    category: str
    size_bytes: int
    
    @property
    def is_multimodal(self) -> bool:
        """Check if file requires multimodal processing (image/video/audio)."""
        return self.category in ["image", "video", "audio"]


@dataclass
class FolderContents:
    """Structured representation of folder contents."""
    folder_path: str
    files: List[FileInfo] = field(default_factory=list)
    
    @property
    def images(self) -> List[FileInfo]:
        return [f for f in self.files if f.category == "image"]
    
    @property
    def videos(self) -> List[FileInfo]:
        return [f for f in self.files if f.category == "video"]
    
    @property
    def audios(self) -> List[FileInfo]:
        return [f for f in self.files if f.category == "audio"]
    
    @property
    def documents(self) -> List[FileInfo]:
        return [f for f in self.files if f.category == "document"]
    
    @property
    def spreadsheets(self) -> List[FileInfo]:
        return [f for f in self.files if f.category == "spreadsheet"]
    
    @property
    def presentations(self) -> List[FileInfo]:
        return [f for f in self.files if f.category == "presentation"]
    
    @property
    def data_files(self) -> List[FileInfo]:
        return [f for f in self.files if f.category == "data"]
    
    @property
    def archives(self) -> List[FileInfo]:
        return [f for f in self.files if f.category == "archive"]
    
    @property
    def other_files(self) -> List[FileInfo]:
        return [f for f in self.files if f.category == "other"]
    
    @property
    def multimodal_files(self) -> List[FileInfo]:
        """Get all files that require multimodal processing."""
        return [f for f in self.files if f.is_multimodal]
    
    @property
    def text_extractable_files(self) -> List[FileInfo]:
        """Get all files that can have text extracted."""
        return [f for f in self.files if f.category in 
                ["document", "spreadsheet", "presentation", "data"]]
    
    def get_summary(self) -> str:
        """Get a summary of folder contents."""
        summary_parts = [f"Folder: {self.folder_path}"]
        summary_parts.append(f"Total files: {len(self.files)}")
        
        categories = {}
        for f in self.files:
            categories[f.category] = categories.get(f.category, 0) + 1
        
        for cat, count in sorted(categories.items()):
            summary_parts.append(f"  - {cat}: {count}")
        
        return "\n".join(summary_parts)


def get_file_category(extension: str) -> str:
    """Determine the category of a file based on its extension."""
    ext = extension.lower()
    
    if ext in IMAGE_EXTENSIONS:
        return "image"
    elif ext in VIDEO_EXTENSIONS:
        return "video"
    elif ext in AUDIO_EXTENSIONS:
        return "audio"
    elif ext in DOCUMENT_EXTENSIONS:
        return "document"
    elif ext in SPREADSHEET_EXTENSIONS:
        return "spreadsheet"
    elif ext in PRESENTATION_EXTENSIONS:
        return "presentation"
    elif ext in DATA_EXTENSIONS:
        return "data"
    elif ext in ARCHIVE_EXTENSIONS:
        return "archive"
    else:
        return "other"


def scan_folder(folder_path: str, recursive: bool = False) -> FolderContents:
    """
    Scan a folder and categorize all files.
    
    Args:
        folder_path: Path to the folder to scan
        recursive: Whether to scan subdirectories recursively
        
    Returns:
        FolderContents object with categorized files
    """
    folder_path = os.path.abspath(folder_path)
    
    if not os.path.exists(folder_path):
        raise FileNotFoundError(f"Folder not found: {folder_path}")
    
    if not os.path.isdir(folder_path):
        raise NotADirectoryError(f"Path is not a directory: {folder_path}")
    
    contents = FolderContents(folder_path=folder_path)
    
    if recursive:
        for root, _, files in os.walk(folder_path):
            for filename in files:
                if filename.startswith("."):  # Skip hidden files
                    continue
                file_path = os.path.join(root, filename)
                _add_file_info(contents, file_path, filename)
    else:
        for filename in os.listdir(folder_path):
            if filename.startswith("."):  # Skip hidden files
                continue
            file_path = os.path.join(folder_path, filename)
            if os.path.isfile(file_path):
                _add_file_info(contents, file_path, filename)
    
    return contents


def _add_file_info(contents: FolderContents, file_path: str, filename: str) -> None:
    """Add file information to FolderContents."""
    _, ext = os.path.splitext(filename)
    category = get_file_category(ext)
    
    try:
        size = os.path.getsize(file_path)
    except OSError:
        size = 0
    
    contents.files.append(FileInfo(
        path=file_path,
        name=filename,
        extension=ext.lower(),
        category=category,
        size_bytes=size
    ))


def _extract_file_content(file_info: FileInfo, max_content_length: int = 200_000) -> Optional[str]:
    """
    Extract text content from a file using existing converters.
    
    Args:
        file_info: FileInfo object for the file
        max_content_length: Maximum length of content to return
        
    Returns:
        Extracted text content or None if extraction failed
    """
    file_path = file_info.path
    ext = file_info.extension.lower()
    
    try:
        parsing_result = None
        
        # Use existing converters from input_handler
        if ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            parsing_result = DocumentConverterResult(title=None, text_content=content)
        
        elif ext in [".json", ".jsonld"]:
            # 跳过所有 JSON 文件，不将其内容提取到任务描述中
            # Long context 文件由 RAG 工具处理
            # 其他 JSON 文件（如 useful_search.json, noise_search.json）也跳过
            return None
        
        elif ext in [".xlsx", ".xls"]:
            # Excel 文件：只显示前 10 行，提供 Python 计算指南
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True)
                
                # 获取所有工作表信息
                sheet_info = []
                for sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                    rows = sheet.max_row
                    cols = sheet.max_column
                    sheet_info.append(f"  - {sheet_name}: {rows} 行 x {cols} 列")
                
                first_sheet = wb[wb.sheetnames[0]]
                total_rows = first_sheet.max_row
                total_cols = first_sheet.max_column
                
                content = f"**[Excel 文件预览 - 仅显示前 10 行]**\n\n"
                content += f"文件路径: `{file_path}`\n"
                content += f"工作表:\n" + "\n".join(sheet_info) + "\n\n"
                
                # 用三个单引号包裹文件内容
                content += f"'''\n"
                
                # 显示前 10 行（所有列）
                content += f"数据预览 (第 1 个工作表: {wb.sheetnames[0]})\n\n"
                content += "|"
                for col_idx in range(1, total_cols + 1):
                    cell = first_sheet.cell(row=1, column=col_idx)
                    cell_value = str(cell.value) if cell.value is not None else ""
                    content += f" {cell_value} |"
                content += "\n|"
                for _ in range(total_cols):
                    content += " --- |"
                content += "\n"
                
                for row_idx in range(2, min(11, total_rows + 1)):  # 前 10 行数据
                    content += "|"
                    for col_idx in range(1, total_cols + 1):
                        cell = first_sheet.cell(row=row_idx, column=col_idx)
                        cell_value = str(cell.value) if cell.value is not None else ""
                        content += f" {cell_value} |"
                    content += "\n"
                
                # 明显的截断符号
                content += f"\n{'='*60}\n"
                content += f"⚠️⚠️⚠️ 【文件已截断】共 {total_rows} 行，仅显示前 10 行 ⚠️⚠️⚠️\n"
                content += f"{'='*60}\n"
                content += f"'''\n\n"
                
                content += f"**🚨 重要提示：这是用户上传的文件，内容已被截断！**\n"
                content += f"- 总行数: {total_rows}\n"
                content += f"- 总列数: {total_cols}\n"
                content += f"- 已显示: 前 10 行\n"
                content += f"- **未显示: 第 11-{total_rows} 行**\n\n"
                content += f"**⚠️ 您必须使用工具读取完整数据！**\n\n"
                filename = os.path.basename(file_path)
                content += f"**方法 1: 使用 Python 代码进行数值计算（推荐）：**\n"
                content += f"```\n"
                content += f"# 步骤 1: 先创建 sandbox\n"
                content += f"create_sandbox()\n"
                content += f"\n"
                content += f"# 步骤 2: 上传文件到 sandbox\n"
                content += f"upload_file_from_local_to_sandbox(sandbox_id=<sandbox_id>, local_file_path='{file_path}')\n"
                content += f"# 上传后文件路径为: /home/user/{filename}\n"
                content += f"\n"
                content += f"# 步骤 3: 在 sandbox 中运行 Python 代码\n"
                content += f"run_python_code(sandbox_id=<sandbox_id>, code_block='''\n"
                content += f"import pandas as pd\n"
                content += f"df = pd.read_excel('/home/user/{filename}')  # 注意：使用 sandbox 中的路径\n"
                content += f"print(df.head())\n"
                content += f"# 计算平均值: print(df['列名'].mean())\n"
                content += f"# 计算总和: print(df['列名'].sum())\n"
                content += f"# 计算差值: print(df['列A'] - df['列B'])\n"
                content += f"''')\n"
                content += f"```\n\n"
                content += f"**⚠️ 重要：在 sandbox 中运行代码时，文件路径是 `/home/user/{filename}`，不是本地路径！**\n\n"
                content += f"**方法 2: 使用工具读取指定行：**\n"
                content += f"- `read_excel_rows(file_path='{file_path}', start_row=N, end_row=M)`\n"
                content += f"- `search_in_file(file_path='{file_path}', keyword='关键词')`\n"
                
                parsing_result = DocumentConverterResult(title=None, text_content=content)
            except Exception as e:
                # 回退到完整转换
                parsing_result = XlsxConverter(local_path=file_path)
        
        elif ext == ".pdf":
            # Extract PDF content - only first page for long PDFs
            if HAS_PDFMINER:
                from pdfminer.pdfpage import PDFPage
                
                # Get page count first
                with open(file_path, 'rb') as f:
                    pages = list(PDFPage.get_pages(f))
                    total_pages = len(pages)
                
                # 对于单页 PDF，直接提取全部内容（用三引号包裹）
                if total_pages == 1:
                    pdf_text = pdfminer.high_level.extract_text(file_path)
                    content = f"**[PDF 文档 - 共 1 页]**\n\n"
                    content += f"文件路径: `{file_path}`\n\n"
                    content += f"'''\n{pdf_text}\n'''\n"
                    parsing_result = DocumentConverterResult(title=None, text_content=content)
                else:
                    # 对于多页 PDF，只提取第一页内容
                    first_page_content = pdfminer.high_level.extract_text(
                        file_path, 
                        page_numbers=[0]  # 只提取第一页（索引从0开始）
                    )
                    
                    content = f"**[PDF 文档预览 - 仅显示第 1 页，共 {total_pages} 页]**\n\n"
                    content += f"文件路径: `{file_path}`\n\n"
                    
                    # 用三引号包裹文件内容
                    content += f"'''\n"
                    content += first_page_content
                    
                    # 明显的截断符号
                    content += f"\n\n{'='*60}\n"
                    content += f"⚠️⚠️⚠️ 【PDF 文件已截断】共 {total_pages} 页，仅显示第 1 页 ⚠️⚠️⚠️\n"
                    content += f"{'='*60}\n"
                    content += f"'''\n\n"
                    
                    content += f"**🚨 重要提示：这是用户上传的 PDF 文件，内容已被截断！**\n"
                    content += f"- 总页数: {total_pages} 页\n"
                    content += f"- 已显示: 第 1 页\n"
                    content += f"- **未显示: 第 2-{total_pages} 页**\n\n"
                    content += f"**⚠️ 您必须使用工具读取完整内容！**\n\n"
                    content += f"**可用工具：**\n"
                    content += f"1. `read_pdf_pages(file_path='{file_path}', start_page=N, end_page=M)` - 读取指定页\n"
                    content += f"   - 例如：读取第2-5页: `read_pdf_pages(file_path='{file_path}', start_page=2, end_page=5)`\n"
                    content += f"2. `search_in_file(file_path='{file_path}', keyword='关键词')` - 搜索关键词\n"
                    content += f"3. `get_file_info(file_path='{file_path}')` - 获取文档结构\n"
                    
                    parsing_result = DocumentConverterResult(title=None, text_content=content)
        
        elif ext in [".docx", ".doc"]:
            parsing_result = DocxConverter(local_path=file_path)
        
        elif ext in [".html", ".htm"]:
            parsing_result = HtmlConverter(local_path=file_path)
        
        elif ext in [".pptx", ".ppt"]:
            parsing_result = PptxConverter(local_path=file_path)
        
        elif ext == ".zip":
            parsing_result = ZipConverter(local_path=file_path)
        
        elif ext == ".csv":
            # CSV 文件：只显示前 10 行，提供 Python 计算指南
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                lines = f.readlines()
            
            total_rows = len(lines)
            
            # 解析表头获取列名
            header = lines[0].rstrip('\n') if lines else ""
            columns = header.split(',')
            col_count = len(columns)
            
            content = f"**[CSV 文件预览 - 仅显示前 10 行]**\n\n"
            content += f"文件路径: `{file_path}`\n"
            content += f"总行数: {total_rows} 行\n"
            content += f"总列数: {col_count} 列\n\n"
            
            # 用三引号包裹文件内容
            content += f"'''\n"
            
            # 显示前 10 行（所有列）- 使用 Markdown 表格格式
            content += f"数据预览:\n\n"
            
            # 表头
            content += "|"
            for col in columns:
                content += f" {col.strip()} |"
            content += "\n|"
            for _ in columns:
                content += " --- |"
            content += "\n"
            
            # 数据行（前 10 行，不包括表头）
            for i, line in enumerate(lines[1:11]):  # 跳过表头，取前 10 行数据
                cols = line.rstrip('\n').split(',')
                content += "|"
                for col in cols:
                    content += f" {col.strip()} |"
                content += "\n"
            
            # 明显的截断符号
            content += f"\n{'='*60}\n"
            content += f"⚠️⚠️⚠️ 【CSV 文件已截断】共 {total_rows} 行，仅显示前 10 行 ⚠️⚠️⚠️\n"
            content += f"{'='*60}\n"
            content += f"'''\n\n"
            
            content += f"**🚨 重要提示：这是用户上传的 CSV 文件，内容已被截断！**\n"
            content += f"- 总行数: {total_rows}\n"
            content += f"- 总列数: {col_count}\n"
            content += f"- 列名: {', '.join(columns[:10])}"
            if col_count > 10:
                content += f" ... (共 {col_count} 列)"
            content += f"\n"
            content += f"- **未显示: 第 11-{total_rows} 行**\n\n"
            content += f"**⚠️ 您必须使用工具读取完整数据！**\n\n"
            filename = os.path.basename(file_path)
            content += f"**方法 1: 使用 Python 代码进行数值计算（推荐）：**\n"
            content += f"```\n"
            content += f"# 步骤 1: 先创建 sandbox\n"
            content += f"create_sandbox()\n"
            content += f"\n"
            content += f"# 步骤 2: 上传文件到 sandbox\n"
            content += f"upload_file_from_local_to_sandbox(sandbox_id=<sandbox_id>, local_file_path='{file_path}')\n"
            content += f"# 上传后文件路径为: /home/user/{filename}\n"
            content += f"\n"
            content += f"# 步骤 3: 在 sandbox 中运行 Python 代码\n"
            content += f"run_python_code(sandbox_id=<sandbox_id>, code_block='''\n"
            content += f"import pandas as pd\n"
            content += f"df = pd.read_csv('/home/user/{filename}')  # 注意：使用 sandbox 中的路径\n"
            content += f"print(df.head())\n"
            content += f"# 计算平均值: print(df['列名'].mean())\n"
            content += f"# 计算总和: print(df['列名'].sum())\n"
            content += f"# 分组统计: print(df.groupby('分组列')['数值列'].mean())\n"
            content += f"''')\n"
            content += f"```\n\n"
            content += f"**⚠️ 重要：在 sandbox 中运行代码时，文件路径是 `/home/user/{filename}`，不是本地路径！**\n\n"
            content += f"**方法 2: 使用工具读取指定行：**\n"
            content += f"- `read_excel_rows(file_path='{file_path}', start_row=N, end_row=M)`\n"
            content += f"- `search_in_file(file_path='{file_path}', keyword='关键词')`\n"
            
            parsing_result = DocumentConverterResult(title=None, text_content=content)
        
        # Try MarkItDown as fallback for other file types
        if parsing_result is None and HAS_MARKITDOWN:
            try:
                md = MarkItDown(enable_plugins=True)
                parsing_result = md.convert(file_path)
            except Exception:
                pass
        
        # Extract content from result
        if parsing_result:
            content = parsing_result.text_content
            if content and len(content) > max_content_length:
                content = content[:max_content_length] + "\n... [Content truncated]"
            return content
        
    except Exception as e:
        return f"[Error extracting content: {str(e)}]"
    
    return None


def _get_image_info(file_info: FileInfo) -> str:
    """Get image information string."""
    info_parts = [f"Image file: {file_info.name}"]
    info_parts.append(f"Path: {file_info.path}")
    
    if HAS_PIL:
        try:
            with Image.open(file_info.path) as img:
                width, height = img.size
                info_parts.append(f"Dimensions: {width}x{height} pixels")
                info_parts.append(f"Format: {img.format}")
        except Exception:
            pass
    
    return "\n".join(info_parts)


def _get_video_info(file_info: FileInfo) -> str:
    """Get video information string."""
    info_parts = [f"Video file: {file_info.name}"]
    info_parts.append(f"Path: {file_info.path}")
    
    try:
        from moviepy.editor import VideoFileClip
        clip = VideoFileClip(file_info.path)
        info_parts.append(f"Duration: {clip.duration:.2f} seconds")
        info_parts.append(f"Resolution: {clip.w}x{clip.h}")
        info_parts.append(f"FPS: {clip.fps:.1f}")
        clip.close()
    except Exception:
        try:
            import cv2
            cap = cv2.VideoCapture(file_info.path)
            fps = cap.get(cv2.CAP_PROP_FPS)
            frame_count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
            height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
            if fps > 0:
                duration = frame_count / fps
                info_parts.append(f"Duration: {duration:.2f} seconds")
            info_parts.append(f"Resolution: {width}x{height}")
            info_parts.append(f"FPS: {fps:.1f}")
            cap.release()
        except Exception:
            pass
    
    return "\n".join(info_parts)


def _get_audio_info(file_info: FileInfo) -> str:
    """Get audio information string."""
    info_parts = [f"Audio file: {file_info.name}"]
    info_parts.append(f"Path: {file_info.path}")
    
    ext = file_info.extension.lower()
    
    if ext == ".wav":
        try:
            import wave
            with wave.open(file_info.path, "rb") as audio_file:
                duration = audio_file.getnframes() / float(audio_file.getframerate())
                sample_rate = audio_file.getframerate()
                channels = audio_file.getnchannels()
                info_parts.append(f"Duration: {duration:.2f} seconds")
                info_parts.append(f"Sample rate: {sample_rate} Hz")
                info_parts.append(f"Channels: {channels}")
        except Exception:
            pass
    else:
        try:
            from mutagen import File as MutagenFile
            audio = MutagenFile(file_info.path)
            if audio and hasattr(audio, "info") and hasattr(audio.info, "length"):
                info_parts.append(f"Duration: {audio.info.length:.2f} seconds")
                if hasattr(audio.info, "sample_rate"):
                    info_parts.append(f"Sample rate: {audio.info.sample_rate} Hz")
        except Exception:
            pass
    
    return "\n".join(info_parts)


def process_folder_for_task(
    folder_path: str,
    query: str,
    recursive: bool = False,
    include_file_contents: bool = True,
    max_content_length: int = 200_000
) -> Tuple[str, str, List[str]]:
    """
    Process all files in a folder and prepare task description for LLM.
    
    This function:
    1. Scans the folder and categorizes all files
    2. Extracts text content from documents, spreadsheets, etc.
    3. Prepares multimodal file information (images, videos, audio)
    4. Generates a comprehensive task description with tool usage guidance
    
    Args:
        folder_path: Path to the folder to process
        query: The user's query/question about the folder contents
        recursive: Whether to scan subdirectories recursively
        include_file_contents: Whether to include extracted file contents
        max_content_length: Maximum length of content per file
        
    Returns:
        Tuple of:
        - task_content: Full content string for LLM (includes file contents)
        - task_description: Task description with tool guidance
        - multimodal_files: List of paths to multimodal files (images, videos, audio)
    """
    # Scan folder
    contents = scan_folder(folder_path, recursive=recursive)
    
    # Build task description
    task_parts = []
    task_parts.append(f"# Task\n\n{query}\n")
    
    # Add folder summary
    task_parts.append(f"\n## Folder Contents Summary\n\n{contents.get_summary()}\n")
    
    # Process text-extractable files (excluding long_context.json files which use RAG)
    local_doc_files = [f for f in contents.text_extractable_files 
                       if "long_context" not in f.name.lower()]
    
    if include_file_contents and local_doc_files:
        task_parts.append("\n## Document Contents (LOCAL FILES - DIRECTLY PROVIDED)\n")
        task_parts.append("\n**⚠️ IMPORTANT: The following document contents are DIRECTLY PROVIDED in this prompt.**")
        task_parts.append("**You MUST cite these files using their file names when referencing their content.**")
        task_parts.append("**Citation format: [filename.ext] or [filename.ext, section/page]**\n")
        
        for file_info in local_doc_files:
            content = _extract_file_content(file_info, max_content_length)
            if content:
                task_parts.append(f"\n### {file_info.name} ⭐ **LOCAL FILE - CITE AS [{file_info.name}]**\n")
                task_parts.append(f"<file path=\"{file_info.path}\" citation=\"[{file_info.name}]\">\n{content}\n</file>\n")
        
        task_parts.append("\n---")
        task_parts.append("**CITATION REMINDER**: When using information from the above files, cite them as:")
        task_parts.append("- For PPT: [filename.pptx, Slide N] or [filename.pptx]")
        task_parts.append("- For PDF: [filename.pdf, Page N] or [filename.pdf]")
        task_parts.append("- For other docs: [filename.ext]\n")
    
    # Process multimodal files
    multimodal_files = []
    
    # Images
    if contents.images:
        task_parts.append("\n## Image Files\n")
        task_parts.append("\nThe following image files are available for analysis:\n")
        
        for file_info in contents.images:
            multimodal_files.append(file_info.path)
            task_parts.append(f"\n### {file_info.name}\n")
            task_parts.append(_get_image_info(file_info))
        
        task_parts.append("\n\n**IMPORTANT**: Use the 'vision_understanding_advanced' tool to analyze these images.")
        task_parts.append("This tool provides multi-turn verification, confidence scoring, and cross-validation.")
        task_parts.append("Recommended approach:")
        task_parts.append("1. Call vision_understanding_advanced with a specific question about each image")
        task_parts.append("2. Review the confidence score and metadata")
        task_parts.append("3. If confidence < 0.75, use follow-up analysis or web search for verification\n")
    
    # Videos
    if contents.videos:
        task_parts.append("\n## Video Files\n")
        task_parts.append("\nThe following video files are available for analysis:\n")
        
        for file_info in contents.videos:
            multimodal_files.append(file_info.path)
            task_parts.append(f"\n### {file_info.name}\n")
            task_parts.append(_get_video_info(file_info))
        
        task_parts.append("\n\n**IMPORTANT**: Use the 'video_understanding_advanced' tool to analyze these videos.")
        task_parts.append("Recommendation:")
        task_parts.append("- Use enable_verification=true for detailed action/scene analysis")
        task_parts.append("- For quick preview, use 'video_quick_analysis' tool")
        task_parts.append("- To analyze specific time ranges, use 'video_temporal_qa' with start_time and end_time")
        task_parts.append("- To extract key moments/frames, use 'video_extract_keyframes' tool\n")
    
    # Audio
    if contents.audios:
        task_parts.append("\n## Audio Files\n")
        task_parts.append("\nThe following audio files are available for analysis:\n")
        
        for file_info in contents.audios:
            multimodal_files.append(file_info.path)
            task_parts.append(f"\n### {file_info.name}\n")
            task_parts.append(_get_audio_info(file_info))
        
        task_parts.append("\n\n**IMPORTANT**: Use the 'audio_understanding_advanced' tool to analyze these audio files.")
        task_parts.append("Recommendation:")
        task_parts.append("- Use enable_verification=true for critical transcriptions")
        task_parts.append("- For quick transcription, use 'audio_quick_transcription' tool")
        task_parts.append("- To answer specific questions about the audio, use 'audio_question_answering_enhanced'\n")
    
    # Long context files (RAG)
    # Check for .db files (pre-built embedding databases) - these have highest priority
    db_files = [f for f in contents.other_files if f.name.endswith('.chunks.db')]
    
    # Also check for .json files as fallback
    long_context_files = [f for f in contents.data_files if "long_context" in f.name.lower()]
    
    if db_files or long_context_files:
        task_parts.append("\n## Long Context Documents (RAG - 任务专属知识库)\n")
        task_parts.append("\n**🔍 什么是 Long Context？**")
        task_parts.append("Long Context 是我们**提前为这个任务检索的网页资料**，是针对当前任务的**专属知识库**。")
        task_parts.append("这些资料包含了与任务相关的背景信息、参考数据和专业知识。\n")
        task_parts.append("**⚠️ 重要提示：**")
        task_parts.append("- Long Context 与上面的用户上传文件（PPT、PDF等）是**不同的来源**")
        task_parts.append("- 用户上传文件：用户直接提供的原始资料")
        task_parts.append("- Long Context：我们预先检索的补充参考资料\n")
        task_parts.append("**🚀 强烈建议：尽可能多地使用 RAG 工具从 Long Context 中检索有效信息！**")
        task_parts.append("这些资料是专门为当前任务准备的，可能包含解决问题的关键信息。\n")
        
        # If there are pre-built .db files, use them directly (highest priority)
        recommended_db_path = None
        if db_files:
            # Sort db files by size (smaller first, as they are likely sampled versions)
            db_files_sorted = sorted(db_files, key=lambda f: f.size_bytes)
            recommended_db = db_files_sorted[0]
            recommended_db_path = recommended_db.path
            
            task_parts.append(f"\n### {recommended_db.name} ⭐ **RECOMMENDED - USE THIS DATABASE**\n")
            task_parts.append(f"Path: {recommended_db.path}\n")
            task_parts.append(f"Size: {recommended_db.size_bytes / 1024:.1f} KB\n")
            task_parts.append(f"**This is a pre-built embedding database. Use this path directly with RAG tools.**\n")
            task_parts.append(f"**⚠️ IMPORTANT: Always use this file path when calling RAG tools: {recommended_db.path}**\n")
            
            # List other db files if any
            for db_file in db_files_sorted[1:]:
                task_parts.append(f"\n### {db_file.name} (Alternative Database)\n")
                task_parts.append(f"Path: {db_file.path}\n")
                task_parts.append(f"Size: {db_file.size_bytes / 1024:.1f} KB\n")
        
        # List json files for reference (but recommend using db files)
        if long_context_files:
            for file_info in long_context_files:
                # Check if this json file has a corresponding db file
                db_path = file_info.path + ".chunks.db"
                has_db = os.path.exists(db_path)
                
                task_parts.append(f"\n### {file_info.name}")
                if has_db and not recommended_db_path:
                    # If no standalone db was found, but this json has a db, recommend it
                    recommended_db_path = db_path
                    task_parts.append(" ⭐ **HAS PRE-BUILT DATABASE**")
                task_parts.append(f"\n")
                task_parts.append(f"Path: {file_info.path}\n")
                task_parts.append(f"Size: {file_info.size_bytes / 1024:.1f} KB\n")
                if has_db:
                    db_size = os.path.getsize(db_path)
                    task_parts.append(f"**Pre-built embedding database**: {db_path} ({db_size / 1024:.1f} KB)\n")
                elif recommended_db_path:
                    task_parts.append(f"*Note: Use the recommended database file instead.*\n")
        
        task_parts.append("\n**HOW TO USE RAG TOOLS**:")
        task_parts.append("- `rag_search`: Semantic search to find relevant passages")
        task_parts.append("- `rag_get_context`: Get concatenated context for answering questions")
        task_parts.append("- `rag_document_stats`: Get document collection statistics")
        task_parts.append("\n**⚠️ IMPORTANT DISTINCTION**:")
        task_parts.append("- **LOCAL FILES** (PPT, PDF above): Content is ALREADY in this prompt. Cite as [filename.ext]")
        task_parts.append("- **RAG DOCUMENTS**: Need to be searched. Cite as [long_context: \"title\", chunk N]")
        task_parts.append("\n**You should use BOTH sources** - local files for primary content, RAG for supplementary research.")
        if recommended_db_path:
            task_parts.append(f"\n**⚠️ CRITICAL: When calling RAG tools, use this json_path: {recommended_db_path}**")
            task_parts.append("This database has pre-built embeddings and will load instantly without regenerating embeddings.\n")
    
    # Other files
    other_files = [f for f in contents.other_files if "long_context" not in f.name.lower()]
    if other_files:
        task_parts.append("\n## Other Files\n")
        for file_info in other_files:
            task_parts.append(f"- {file_info.name} ({file_info.extension})\n")
    
    # Add output format requirement
    use_cn_prompt = os.environ.get("USE_CN_PROMPT", "0")
    if use_cn_prompt == "1":
        task_parts.append("\n请通过任务分解和MCP工具调用来解决给定的问题。**请生成完整的报告内容，不需要使用 \\boxed{} 包裹。**")
    else:
        task_parts.append("\nYou should follow the format instruction in the request strictly. Generate the complete report content without wrapping it in \\boxed{}.")
    
    task_content = "\n".join(task_parts)
    task_description = task_content
    
    return task_content, task_description, multimodal_files


def process_folder_batch(
    folder_paths: List[str],
    query: str,
    recursive: bool = False,
    include_file_contents: bool = True,
    max_content_length: int = 200_000
) -> List[Tuple[str, str, str, List[str]]]:
    """
    Process multiple folders in batch.
    
    Args:
        folder_paths: List of folder paths to process
        query: The user's query/question about the folder contents
        recursive: Whether to scan subdirectories recursively
        include_file_contents: Whether to include extracted file contents
        max_content_length: Maximum length of content per file
        
    Returns:
        List of tuples, each containing:
        - folder_path: The original folder path
        - task_content: Full content string for LLM
        - task_description: Task description with tool guidance
        - multimodal_files: List of paths to multimodal files
    """
    results = []
    
    for folder_path in folder_paths:
        try:
            task_content, task_description, multimodal_files = process_folder_for_task(
                folder_path=folder_path,
                query=query,
                recursive=recursive,
                include_file_contents=include_file_contents,
                max_content_length=max_content_length
            )
            results.append((folder_path, task_content, task_description, multimodal_files))
        except Exception as e:
            error_msg = f"Error processing folder {folder_path}: {str(e)}"
            results.append((folder_path, error_msg, error_msg, []))
    
    return results
